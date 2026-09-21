#!/usr/bin/env python3
"""veri/adaylar.json → cikti/musteri-adaylari.html + .xlsx

HTML tek dosyadır (dış bağımlılık yok), Excel openpyxl ile yazılır.
Kullanım:  python3 rapor.py
"""

from __future__ import annotations

import html
import json
import pathlib
import sys
from collections import Counter

KOK = pathlib.Path(__file__).resolve().parent
GIRDI = KOK / "veri" / "adaylar.json"
CIKTI = KOK / "cikti"

BILESEN_ADLARI = {
    "urun_uyumu": ("Ürün uyumu", 30),
    "talep": ("Talep sinyali", 20),
    "dijital": ("Dijital olgunluk", 20),
    "erisim": ("Erişilebilirlik", 20),
    "konum": ("Konum", 10),
}


# ── HTML ──────────────────────────────────────────────────────────────────
SAYFA = """<!DOCTYPE html>
<html lang="tr">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Aday Müşteriler</title>
<style>
  :root {{
    --kagit:#faf7f2; --yuzey:#ffffff; --murekkep:#1b1714; --soluk:#6d635a;
    --cizgi:#e7ded1; --vurgu:#8a5a2b; --a:#2f7d55; --b:#9a7b18; --c:#8a6a5c;
    --golge:0 1px 2px rgba(27,23,20,.06), 0 8px 24px rgba(27,23,20,.05);
  }}
  @media (prefers-color-scheme: dark) {{
    :root:not([data-theme="light"]) {{
      --kagit:#14110f; --yuzey:#1c1815; --murekkep:#f2ece4; --soluk:#a2978a;
      --cizgi:#2e2822; --vurgu:#d2a06a; --a:#6fc194; --b:#d8bc5c; --c:#c2a08f;
      --golge:0 1px 2px rgba(0,0,0,.4), 0 8px 24px rgba(0,0,0,.3);
    }}
  }}
  :root[data-theme="dark"] {{
    --kagit:#14110f; --yuzey:#1c1815; --murekkep:#f2ece4; --soluk:#a2978a;
    --cizgi:#2e2822; --vurgu:#d2a06a; --a:#6fc194; --b:#d8bc5c; --c:#c2a08f;
  }}
  * {{ box-sizing:border-box; }}
  body {{
    margin:0; background:var(--kagit); color:var(--murekkep);
    font:16px/1.55 ui-sans-serif,-apple-system,"Segoe UI",Inter,system-ui,sans-serif;
    -webkit-font-smoothing:antialiased;
  }}
  .sarmal {{ max-width:1180px; margin:0 auto; padding:48px 16px 96px; }}
  header p {{ color:var(--soluk); max-width:62ch; }}
  h1 {{ font-size:clamp(28px,4vw,42px); line-height:1.1; margin:0 0 8px; letter-spacing:-.02em; }}
  .etiket {{ font-size:12px; letter-spacing:.16em; text-transform:uppercase; color:var(--vurgu); margin:0 0 10px; }}
  .kpi {{ display:grid; gap:12px; grid-template-columns:repeat(auto-fit,minmax(150px,1fr)); margin:32px 0; }}
  .kutu {{ background:var(--yuzey); border:1px solid var(--cizgi); border-radius:14px; padding:16px 18px; box-shadow:var(--golge); }}
  .kutu b {{ display:block; font-size:28px; font-weight:650; letter-spacing:-.02em; }}
  .kutu span {{ font-size:13px; color:var(--soluk); }}
  .araclar {{ display:flex; flex-wrap:wrap; gap:8px; align-items:center; margin:24px 0 14px; }}
  input[type=search], select {{
    font:inherit; font-size:14px; padding:9px 12px; border-radius:10px;
    border:1px solid var(--cizgi); background:var(--yuzey); color:var(--murekkep); min-width:0;
  }}
  input[type=search] {{ flex:1 1 220px; }}
  .cip {{ border:1px solid var(--cizgi); background:var(--yuzey); color:var(--soluk);
    border-radius:999px; padding:8px 14px; font-size:13px; cursor:pointer; }}
  .cip[aria-pressed="true"] {{ background:var(--murekkep); color:var(--kagit); border-color:var(--murekkep); }}
  .tablo-kap {{ overflow-x:auto; border:1px solid var(--cizgi); border-radius:14px; background:var(--yuzey); box-shadow:var(--golge); }}
  table {{ border-collapse:collapse; width:100%; font-size:14px; min-width:860px; }}
  th, td {{ padding:12px 14px; text-align:left; border-bottom:1px solid var(--cizgi); vertical-align:top; }}
  th {{ position:sticky; top:0; background:var(--yuzey); font-size:12px; letter-spacing:.06em;
       text-transform:uppercase; color:var(--soluk); cursor:pointer; white-space:nowrap; z-index:1; }}
  tr:last-child td {{ border-bottom:0; }}
  tbody tr:hover {{ background:color-mix(in oklab, var(--vurgu) 6%, transparent); }}
  .isim {{ font-weight:600; }}
  .alt {{ color:var(--soluk); font-size:12.5px; }}
  a {{ color:var(--vurgu); text-decoration:none; }}
  a:hover {{ text-decoration:underline; }}
  .rozet {{ display:inline-block; min-width:22px; text-align:center; font-weight:700; font-size:12px;
    padding:3px 8px; border-radius:999px; color:#fff; }}
  .rozet.A {{ background:var(--a); }} .rozet.B {{ background:var(--b); }} .rozet.C {{ background:var(--c); }}
  .puan {{ display:flex; align-items:center; gap:8px; }}
  .cubuk {{ width:62px; height:6px; border-radius:3px; background:var(--cizgi); overflow:hidden; }}
  .cubuk i {{ display:block; height:100%; background:var(--vurgu); }}
  .yok {{ color:var(--soluk); }}
  footer {{ margin-top:48px; color:var(--soluk); font-size:13.5px; }}
  footer h2 {{ color:var(--murekkep); font-size:18px; margin:28px 0 8px; }}
  footer li {{ margin:4px 0; }}
  @media (max-width:640px) {{ .sarmal {{ padding:32px 16px 72px; }} }}
</style>
</head>
<body>
<div class="sarmal">
  <header>
    <p class="etiket">YSMN · toptan satış geliştirme</p>
    <h1>Aday müşteri listesi</h1>
    <p>Türkiye'de çanta satan perakende noktaları — Google Haritalar'dan taranıp
       kendi sitelerinden iletişim bilgisi toplanarak puanlandı. Puan, satın alma
       olasılığını değil <em>ulaşılabilir ve uygun</em> olma derecesini ölçer.</p>
  </header>

  <section class="kpi">
    <div class="kutu"><b>{adet}</b><span>aday nokta</span></div>
    <div class="kutu"><b>{a_adet}</b><span>A kademesi (78+ puan)</span></div>
    <div class="kutu"><b>{eposta_adet}</b><span>e-postası bulunan</span></div>
    <div class="kutu"><b>{sehir_adet}</b><span>şehir</span></div>
    <div class="kutu"><b>{maliyet}</b><span>veri maliyeti (USD)</span></div>
  </section>

  <div class="araclar">
    <input type="search" id="ara" placeholder="İsim, şehir, alan adı ara…" aria-label="Ara">
    <button class="cip" data-kademe="hepsi" aria-pressed="true">Hepsi</button>
    <button class="cip" data-kademe="A" aria-pressed="false">A</button>
    <button class="cip" data-kademe="B" aria-pressed="false">B</button>
    <button class="cip" data-kademe="C" aria-pressed="false">C</button>
    <select id="sehir" aria-label="Şehir"><option value="">Tüm şehirler</option>{sehir_secenekleri}</select>
    <button class="cip" id="epostali" aria-pressed="false">Yalnız e-postalı</button>
  </div>

  <div class="tablo-kap">
    <table id="tablo">
      <thead><tr>
        <th data-tur="sayi">#</th><th>Firma</th><th data-tur="sayi">Puan</th>
        <th>Şehir</th><th>E-posta</th><th>Telefon</th><th>Web</th><th data-tur="sayi">Yorum</th>
      </tr></thead>
      <tbody>{satirlar}</tbody>
    </table>
  </div>
  <p class="alt" id="sayac"></p>

  <footer>
    <h2>Puan nasıl hesaplandı?</h2>
    <ul>
      <li><b>Ürün uyumu (30)</b> — mağaza kategorisi çantaya ne kadar yakın: çanta/saraciye tam puan, giyim butiği kısmi.</li>
      <li><b>Talep sinyali (20)</b> — Haritalar yorum sayısı ve yıldızı; mağazanın müşteri trafiği vekili.</li>
      <li><b>Dijital olgunluk (20)</b> — web sitesi, e-ticaret izi, Instagram hesabı.</li>
      <li><b>Erişilebilirlik (20)</b> — kurumsal e-posta en yüksek, genel e-posta orta, telefon ve DM düşük.</li>
      <li><b>Konum (10)</b> — sevkiyat ve saha ziyareti kolaylığı; üç şehir kademesi.</li>
    </ul>
    <h2>Kapsam</h2>
    <p>{kapsam}</p>
    <p>Toptancı, imalatçı ve tamirciler listeden elendi; aynı firma birden çok
       şubeyle çıktıysa tek kayda indirildi. Oluşturma: {tarih}.</p>
  </footer>
</div>
<script>
const satirlar = [...document.querySelectorAll('#tablo tbody tr')];
const sayac = document.getElementById('sayac');
let kademe = 'hepsi', yalnizEposta = false;

function suz() {{
  const q = document.getElementById('ara').value.trim().toLowerCase();
  const sehir = document.getElementById('sehir').value;
  let gorunen = 0;
  for (const tr of satirlar) {{
    const uygun =
      (kademe === 'hepsi' || tr.dataset.kademe === kademe) &&
      (!sehir || tr.dataset.sehir === sehir) &&
      (!yalnizEposta || tr.dataset.eposta === '1') &&
      (!q || tr.dataset.arama.includes(q));
    tr.hidden = !uygun;
    if (uygun) gorunen++;
  }}
  sayac.textContent = gorunen + ' / ' + satirlar.length + ' aday gösteriliyor';
}}

document.getElementById('ara').addEventListener('input', suz);
document.getElementById('sehir').addEventListener('change', suz);
document.getElementById('epostali').addEventListener('click', (e) => {{
  yalnizEposta = !yalnizEposta;
  e.currentTarget.setAttribute('aria-pressed', String(yalnizEposta));
  suz();
}});
for (const cip of document.querySelectorAll('.cip[data-kademe]')) {{
  cip.addEventListener('click', () => {{
    kademe = cip.dataset.kademe;
    document.querySelectorAll('.cip[data-kademe]').forEach(d =>
      d.setAttribute('aria-pressed', String(d === cip)));
    suz();
  }});
}}

const govde = document.querySelector('#tablo tbody');
document.querySelectorAll('#tablo th').forEach((th, i) => {{
  const sayi = th.dataset.tur === 'sayi';
  // Sayı sütunları önce büyükten küçüğe, metin sütunları A'dan Z'ye açılır.
  let artan = sayi;
  th.addEventListener('click', () => {{
    artan = !artan;
    [...govde.rows].sort((a, b) => {{
      const x = a.cells[i].dataset.v ?? a.cells[i].textContent;
      const y = b.cells[i].dataset.v ?? b.cells[i].textContent;
      const s = sayi ? (+x) - (+y) : String(x).localeCompare(String(y), 'tr');
      return artan ? s : -s;
    }}).forEach(tr => govde.appendChild(tr));
  }});
}});
suz();
</script>
</body>
</html>
"""


def _kacis(metin: str) -> str:
    return html.escape(str(metin or ""))


def satir_yaz(aday: dict) -> str:
    eposta = aday.get("eposta")
    eposta_hucre = (
        f'<a href="mailto:{_kacis(eposta)}">{_kacis(eposta)}</a>'
        + (f'<div class="alt">{_kacis(aday["eposta_turu"])}</div>' if aday.get("eposta_turu") else "")
        if eposta else '<span class="yok">—</span>'
    )
    site = aday.get("site")
    site_hucre = (
        f'<a href="{_kacis(site)}" target="_blank" rel="noopener">{_kacis(aday.get("alan_adi") or site)}</a>'
        if site else '<span class="yok">—</span>'
    )
    tel = aday.get("telefon")
    tel_hucre = f'<a href="tel:{_kacis(tel).replace(" ", "")}">{_kacis(tel)}</a>' if tel else '<span class="yok">—</span>'
    harita = aday.get("harita")
    ad_alt = " · ".join(x for x in [aday.get("kategori"), aday.get("ilce")] if x)
    arama = " ".join(str(aday.get(k, "")) for k in ("isim", "sehir", "alan_adi", "kategori", "eposta")).lower()

    return f"""<tr data-kademe="{_kacis(aday['kademe'])}" data-sehir="{_kacis(aday.get('sehir'))}"
  data-eposta="{1 if eposta else 0}" data-arama="{_kacis(arama)}">
  <td data-v="{aday['sira']}">{aday['sira']}</td>
  <td><div class="isim">{_kacis(aday['isim'])}</div>
      <div class="alt">{_kacis(ad_alt)}{' · ' if ad_alt and harita else ''}{f'<a href="{_kacis(harita)}" target="_blank" rel="noopener">harita</a>' if harita else ''}</div></td>
  <td data-v="{aday['puan']}"><div class="puan"><span class="rozet {_kacis(aday['kademe'])}">{aday['puan']}</span>
      <span class="cubuk"><i style="width:{aday['puan']}%"></i></span></div></td>
  <td>{_kacis(aday.get('sehir') or '—')}</td>
  <td>{eposta_hucre}</td>
  <td>{tel_hucre}</td>
  <td>{site_hucre}</td>
  <td data-v="{aday.get('yorum', 0)}">{aday.get('yorum', 0)}<div class="alt">★ {aday.get('yildiz') or '—'}</div></td>
</tr>"""


def html_yaz(paket: dict, yol: pathlib.Path) -> None:
    adaylar = paket["adaylar"]
    sehirler = sorted({a.get("sehir") for a in adaylar if a.get("sehir")})
    kapsam = paket.get("kapsam", {})
    sayfa = SAYFA.format(
        adet=len(adaylar),
        a_adet=sum(1 for a in adaylar if a["kademe"] == "A"),
        eposta_adet=sum(1 for a in adaylar if a.get("eposta")),
        sehir_adet=len(sehirler),
        maliyet=f"{paket.get('maliyet_usd', 0):.2f}",
        sehir_secenekleri="".join(f'<option value="{_kacis(s)}">{_kacis(s)}</option>' for s in sehirler),
        satirlar="\n".join(satir_yaz(a) for a in adaylar),
        kapsam=_kacis(
            f"{len(kapsam.get('sehirler', []))} şehir × {len(kapsam.get('sorgular', []))} arama sorgusu; "
            f"{kapsam.get('ham_kayit', 0)} ham kayıt tarandı, eleme sonrası {kapsam.get('eleme_sonrasi', 0)} "
            f"nokta kaldı, en iyi {len(adaylar)} tanesi listelendi."
        ),
        tarih=_kacis(paket.get("olusturma", "")),
    )
    yol.write_text(sayfa, encoding="utf-8")


# ── Excel ─────────────────────────────────────────────────────────────────
def excel_yaz(paket: dict, yol: pathlib.Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    adaylar = paket["adaylar"]
    kitap = Workbook()

    sayfa = kitap.active
    sayfa.title = "Adaylar"
    basliklar = [
        "#", "Kademe", "Puan", "Firma", "Kategori", "Şehir", "İlçe",
        "E-posta", "E-posta türü", "Telefon", "Web sitesi", "Instagram",
        "Yıldız", "Yorum", "Ürün uyumu", "Talep", "Dijital", "Erişim", "Konum",
        "Zincir mi?", "Adres", "Harita",
    ]
    sayfa.append(basliklar)
    baslik_dolgu = PatternFill("solid", fgColor="1B1714")
    for hucre in sayfa[1]:
        hucre.font = Font(bold=True, color="FFFFFF", size=11)
        hucre.fill = baslik_dolgu
        hucre.alignment = Alignment(vertical="center")

    dolgu = {
        "A": PatternFill("solid", fgColor="D8F0E2"),
        "B": PatternFill("solid", fgColor="FBF0CC"),
        "C": PatternFill("solid", fgColor="F1E8E3"),
    }
    for aday in adaylar:
        b = aday["bilesenler"]
        sayfa.append([
            aday["sira"], aday["kademe"], aday["puan"], aday["isim"], aday.get("kategori"),
            aday.get("sehir"), aday.get("ilce"), aday.get("eposta"), aday.get("eposta_turu"),
            aday.get("telefon"), aday.get("site"), aday.get("instagram"),
            aday.get("yildiz"), aday.get("yorum"),
            b["urun_uyumu"], b["talep"], b["dijital"], b["erisim"], b["konum"],
            "evet" if aday.get("zincir") else "", aday.get("adres"), aday.get("harita"),
        ])
        satir = sayfa.max_row
        sayfa.cell(satir, 2).fill = dolgu.get(aday["kademe"], dolgu["C"])
        sayfa.cell(satir, 2).alignment = Alignment(horizontal="center")
        sayfa.cell(satir, 3).font = Font(bold=True)
        if aday.get("eposta"):
            hucre = sayfa.cell(satir, 8)
            hucre.hyperlink = f"mailto:{aday['eposta']}"
            hucre.font = Font(color="8A5A2B", underline="single")
        if aday.get("site"):
            hucre = sayfa.cell(satir, 11)
            hucre.hyperlink = aday["site"]
            hucre.font = Font(color="8A5A2B", underline="single")

    genislikler = [5, 8, 7, 34, 22, 14, 16, 30, 13, 17, 30, 24, 8, 8, 12, 8, 9, 9, 8, 10, 42, 22]
    for i, genislik in enumerate(genislikler, 1):
        sayfa.column_dimensions[get_column_letter(i)].width = genislik
    sayfa.freeze_panes = "D2"
    sayfa.auto_filter.ref = f"A1:{get_column_letter(len(basliklar))}{sayfa.max_row}"

    ozet = kitap.create_sheet("Özet")
    kademeler = Counter(a["kademe"] for a in adaylar)
    sehirler = Counter(a.get("sehir") or "—" for a in adaylar)
    satirlar = [
        ("Aday sayısı", len(adaylar)),
        ("E-postası bulunan", sum(1 for a in adaylar if a.get("eposta"))),
        ("Kurumsal e-posta", sum(1 for a in adaylar if a.get("eposta_turu") == "kurumsal")),
        ("Web sitesi olan", sum(1 for a in adaylar if a.get("site"))),
        ("Ortalama puan", round(sum(a["puan"] for a in adaylar) / max(1, len(adaylar)), 1)),
        ("", ""),
        ("A kademesi (78+)", kademeler.get("A", 0)),
        ("B kademesi (62-77)", kademeler.get("B", 0)),
        ("C kademesi (<62)", kademeler.get("C", 0)),
        ("", ""),
        ("Apify maliyeti (USD)", paket.get("maliyet_usd", 0)),
        ("  Google Maps", paket.get("maliyet_dokum", {}).get("google_maps", 0)),
        ("  İletişim taraması", paket.get("maliyet_dokum", {}).get("iletisim", 0)),
        ("Oluşturma", paket.get("olusturma", "")),
        ("", ""),
        ("Şehir dağılımı", ""),
    ]
    satirlar += sorted(sehirler.items(), key=lambda p: -p[1])
    for etiket, deger in satirlar:
        ozet.append([etiket, deger])
    for hucre in ozet["A"]:
        hucre.font = Font(bold=hucre.value in {"Aday sayısı", "Şehir dağılımı", "Apify maliyeti (USD)"})
    ozet.column_dimensions["A"].width = 26
    ozet.column_dimensions["B"].width = 22

    yontem = kitap.create_sheet("Yöntem")
    yontem.append(["Bileşen", "Azami", "Ne ölçüyor"])
    for hucre in yontem[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = baslik_dolgu
    for anahtar, (ad, azami) in BILESEN_ADLARI.items():
        aciklama = {
            "urun_uyumu": "Mağaza kategorisinin çantaya yakınlığı (çanta/saraciye tam, giyim butiği kısmi).",
            "talep": "Google Haritalar yorum sayısı ve yıldızı — müşteri trafiği vekili.",
            "dijital": "Web sitesi, e-ticaret izi ve Instagram hesabı.",
            "erisim": "Kurumsal e-posta > genel e-posta > telefon > DM.",
            "konum": "Sevkiyat ve saha ziyareti kolaylığı (şehir kademesi).",
        }[anahtar]
        yontem.append([ad, azami, aciklama])
    yontem.append([])
    yontem.append(["Kademe A", "78-100", "Önce bunlar aranır: uyum, trafik ve iletişim üçü birden güçlü."])
    yontem.append(["Kademe B", "62-77", "İkinci dalga; genelde e-posta eksik, telefonla açılır."])
    yontem.append(["Kademe C", "0-61", "Uzun vade; kategori uyumu ya da iletişim zayıf."])
    yontem.column_dimensions["A"].width = 20
    yontem.column_dimensions["B"].width = 10
    yontem.column_dimensions["C"].width = 90
    for satir in yontem.iter_rows(min_row=2):
        satir[2].alignment = Alignment(wrap_text=True, vertical="top")

    kitap.save(yol)


def main() -> int:
    global GIRDI, CIKTI
    if len(sys.argv) > 1:
        GIRDI = pathlib.Path(sys.argv[1])
    if len(sys.argv) > 2:
        CIKTI = pathlib.Path(sys.argv[2])
    if not GIRDI.exists():
        print(f"HATA: {GIRDI} yok. Önce: APIFY_TOKEN=... python3 topla.py", file=sys.stderr)
        return 1
    paket = json.loads(GIRDI.read_text(encoding="utf-8"))
    CIKTI.mkdir(parents=True, exist_ok=True)
    html_yolu = CIKTI / "musteri-adaylari.html"
    excel_yolu = CIKTI / "musteri-adaylari.xlsx"
    html_yaz(paket, html_yolu)
    excel_yaz(paket, excel_yolu)
    print(f"✓ {html_yolu}")
    print(f"✓ {excel_yolu}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
